#!/usr/bin/env python3
"""生成两套比赛 PPT：交件方案版 + 现场演示讲解版。"""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

_REPO = Path(__file__).resolve().parent.parent
_OUT = _REPO / "dist"
_SHOTS = _REPO / "docs" / "ppt-shots"
_DESKTOP = Path.home() / "Desktop"

# 深色科技风（避免紫）
BG = RGBColor(0x09, 0x09, 0x0B)
FG = RGBColor(0xEA, 0xF2, 0xFF)
DIM = RGBColor(0x8A, 0xA0, 0xB8)
ACCENT = RGBColor(0x3D, 0xD6, 0xFF)


def _set_run(run, size=18, bold=False, color=FG, font="Microsoft YaHei"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    try:
        run._r.get_or_add_rPr().set(  # type: ignore[attr-defined]
            "{http://schemas.openxmlformats.org/officeDocument/2006/main}eastAsia",
            font,
        )
    except Exception:
        pass


def _fill_bg(slide, prs):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def _title(slide, text, top=0.35, size=28):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12.1), Inches(0.7))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=True, color=ACCENT)
    return box


def _body(slide, lines, top=1.2, size=16, width=12.1):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(width), Inches(5.8))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = line
        _set_run(run, size=size, color=FG if not line.startswith("·") else DIM)
    return box


def _notes(slide, text: str):
    slide.notes_slide.notes_text_frame.text = text


def _new(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _fill_bg(slide, prs)
    return slide


def _shot_slide(prs, title: str, filename: str, caption: str):
    """交件版产品截图页：标题 + 大图 + 一句说明。"""
    s = _new(prs)
    _title(s, title, top=0.25, size=24)
    path = _SHOTS / filename
    if path.is_file():
        # 16:9 工作台截图，尽量铺满
        s.shapes.add_picture(str(path), Inches(0.55), Inches(1.0),
                             width=Inches(12.2))
    else:
        _body(s, [f"（缺少截图 {filename}，请先抓取 docs/ppt-shots/）"],
              top=2.5)
    cap = s.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(12.1), Inches(0.4))
    tf = cap.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = caption
    _set_run(r, size=13, color=DIM)
    _notes(s, caption)
    return s


def build_submission() -> Path:
    """交件用：方案讲解，偏完整。"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    s = _new(prs)
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(2.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "veritas"
    _set_run(r, size=48, bold=True, color=ACCENT)
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "多智能体可信研判与证据链溯源"
    _set_run(r, size=28, bold=True, color=FG)
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "OPC 2026 · AI for Science · 方案说明（交件版）"
    _set_run(r, size=16, color=DIM)
    _notes(s, "交件方案 PPT。现场演示请用另一套《现场讲解版》。")

    s = _new(prs)
    _title(s, "1. 我们解决什么问题")
    _body(s, [
        "痛点：AI 能写结论，却很难让人信结论。",
        "· 结论很自信，但说不清依据出自哪段原文",
        "· 多智能体只是简单投票，冲突被抹平",
        "· 报告生成后改了也不留痕，事后无法审计",
        "",
        "产品目标一句话：",
        "多视角独立判断 → 结论挂证据 → 证据接到原文 → 冲突能处理 → 人能改、改能查",
    ])

    s = _new(prs)
    _title(s, "2. 对照赛题五条硬需求")
    _body(s, [
        "① 多智能体协同：三视角硬隔离独立研判 → 归并 / 质证 / 仲裁 → 少数意见保留",
        "② 证据链自动构建：结论 → 逐字引文 → 原文段落；引文不过关就打回（零 LLM）",
        "③ 置信度量化：强度 × 覆盖 × 一致性，可解释、改完可重算（规则分，非概率）",
        "④ 冲突识别与消解：列冲突、质证、按证据强度裁决，落败意见不删",
        "⑤ 轨迹可视化 + 人改重算：工作台下钻 / 推理图；删证据改权重即时重算并入账本",
        "",
        "加分项：相关 vs 因果标注 · 哈希链防篡改 · MCP 标准接口",
    ], size=15)

    s = _new(prs)
    _title(s, "3. 产品形态：三栏工作台")
    _body(s, [
        "界面：左选案件与资料 · 中写议题并开跑 · 右看判断与核对",
        "路径：选/建案件 → 添加资料 → 写议题 → 开始研判（可看进度 / 可终止）→ 核对",
        "",
        "工作台里能做的事：",
        "· 三级下钻：结论 → 引文 → 原文",
        "· 看冲突、质证、裁决与少数意见；主席层给一句话正面回答",
        "· 改结论 / 调权重 / 删证据 → 置信度与综合意见重算并入账本",
        "· 核验账本；「改一字即红」篡改演示",
        "· 接入模型（Base URL + Key + 多模型）；无 Key 可走演示桩",
        "· 议题未齐（缺期间 / 主体等）机械回读早停，不瞎填数字",
        "",
        "演示案例同一引擎：固态电池 · 数据中台 · 谣言vs官方 · 恒兴真公告评测",
    ], size=14)

    # 交件无视频时，截图是评委「看见产品」的主通道——对齐演示三动作
    _shot_slide(
        prs, "产品演示① 工作台总览",
        "01-overview.png",
        "三栏：左案件资料 · 中议题开跑 · 右判断核对；右上角账本可核验",
    )
    _shot_slide(
        prs, "产品演示② 结论 → 引文 → 原文",
        "02-drilldown.png",
        "点「查看原文」：引文对得上资料段落，不是模型转述",
    )
    _shot_slide(
        prs, "产品演示③ 编数字亮黄牌",
        "03-yellow-flags.png",
        "结论里写入资料没有的「99%」→ 程序复核黄牌（来源外数字）",
    )
    _shot_slide(
        prs, "产品演示④ 改一字即红",
        "04-tamper-red.png",
        "篡改演示：诚实账本绿，内存改一处立刻红（磁盘不动）",
    )

    s = _new(prs)
    _title(s, "4. 为什么可信：模型主张，程序复算")
    _body(s, [
        "零 LLM 机械闸（可复现、可测试）：",
        "· 引文逐字核对 · 数字复算 / 溯源 · 归并守恒 · 质证与仲裁引用核对",
        "· 议题前提 · 时序 / as-of · 事实对撞台账 · 注入扫描",
        "· 答案槽位契约：缺期间/主体早停；标题误报「资料不足」可按多数结论纠偏",
        "",
        "协调不是投票：独立写 → 对立就吵 → 程序算证据强度 → 少数意见保留",
        "主席汇总：正面回答议题；漏点程序补全；矛盾「不纳入」说明会清掉",
        "人是最高权限：改动进哈希链，可审计，也可一键重置",
        "",
        "诚实边界（答辩必讲）：",
        "· 置信度是规则融合分，不是校准概率",
        "· 因果是标注加权，不是完整因果推断引擎",
        "· 哈希链是过程账本，不是数字证书 PKI",
    ], size=14)

    s = _new(prs)
    _title(s, "4b. 评测与真公告实测")
    _body(s, [
        "① 标注集三指标（零 LLM，可复现）· eval/benchmark.jsonl + scripts/run_eval.py",
        "· BM25 Recall@K · 冲突召回 · 编造检出（假引文必拒）",
        "",
        "② 真公告实测（工作台真模型）· 恒兴新材 8 份巨潮公告 / 10 问",
        "· 覆盖产销、分红、异动、年报/中报、回购对撞、工商、股权、负向资料不足",
        "· 约 9.5/10；成绩单 eval/hengxing-10q-scorecard.md；答辩卡 hengxing-defense.md",
        "· 现场可切「恒兴新材-评测」看 Q8：营收 / 归母净利 / 同比三数齐",
        "",
        "已知短板不藏：BM25 无语义；难 PDF 默认轻量抽字，Docling 可插拔非默认",
        "数字以现场重跑 / 工作台样例为准，不背死分数",
    ], size=14)

    s = _new(prs)
    _title(s, "4c. 难 PDF / 版式：诚实边界")
    _body(s, [
        "① 文字层 PDF：轻量抽取（pypdf），按页拼文本再切段",
        "② 扫描件 / 纯图 PDF：分诊披露「未解析」，不假装读过",
        "③ 开源版式引擎（Docling）= 可选摄取后端，不塞进分析师人格",
        "   开启：VERITAS_PDF_BACKEND=docling（需另装 requirements-docling）",
        "   产物：派生信源、低可信先验，表格尽量 Markdown 入卷",
        "",
        "跨页无表头、横版合并单元格：业界也未满分——我们用",
        "「读不清就披露 + 入卷可降权 + 闸可验」守住可信主轴",
    ], size=15)

    s = _new(prs)
    _title(s, "5. 可复用与可对接（答辩重点）")
    _body(s, [
        "① 编队可配置：换 roster YAML 即换角色，协调流水线不变",
        "② 模型可替换：同一网关 / 多模型轮转；单模型也能跑完全流程",
        "③ 标准接口：MCP 六工具——读报告、验链、跑研判、外部结论入卷",
        "",
        "我们卖的是可信研判中枢，不是又一个封闭聊天机器人。",
        "第三方按同一证据契约接入；假引文仍过机械闸。",
    ])

    s = _new(prs)
    _title(s, "6. 交付与验证（贴合赛制）")
    _body(s, [
        "评审权重：原型实测 70% + 方案答辩 30%",
        "交件通道：智汇科学城小程序 → 方案 PPT + 方案文档 + 源码包（本轮不要求演示视频）",
        "",
        "· 源码 + 一键部署（Windows / macOS）；断网可用离线演示桩兜底",
        "· 自动化测试全桩零密钥全绿；工作台可现场三动作实测",
        "· 方案文档：docs/proposal.md；现场另备讲解 PPT + 工作台 live",
        "· 真公告样例与成绩单入库，答辩可指屏",
        "",
        "收束三句话：",
        "挂得住原文 · 吵得清并保留少数意见 · 改得明且入账本",
    ], size=15)

    s = _new(prs)
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(11.5), Inches(2))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "谢谢各位老师"
    _set_run(r, size=36, bold=True, color=ACCENT)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "veritas · 可信不是说得更像，而是当场能查"
    _set_run(r, size=16, color=DIM)

    _OUT.mkdir(parents=True, exist_ok=True)
    path = _OUT / "veritas-交件方案.pptx"
    prs.save(path)
    return path


def build_onsite() -> Path:
    """现场讲解用：短、跟演示动作对齐，备注里写口播。"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    s = _new(prs)
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.5), Inches(3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "veritas 现场演示"
    _set_run(r, size=40, bold=True, color=ACCENT)
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "只做三下：摸回原文 · 编数字亮灯 · 改一字即红"
    _set_run(r, size=22, color=FG)
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "工作台 http://127.0.0.1:8765/  ·  流程案：固态电池  ·  被问真公告：恒兴新材"
    _set_run(r, size=14, color=DIM)
    _notes(s, "开场 25 秒。背一句：可信不是模型说得更像，而是三件事当场能查。")

    s = _new(prs)
    _title(s, "开场：我们不比谁写得更长")
    _body(s, [
        "多个角色先各自独立写结论，互相看不见。",
        "意见对立 → 质证 → 裁决；落败意见保留，不删。",
        "最后有一份综合意见，正面回答议题。",
        "",
        "今天只演示三个动作（流程用固态电池）。",
        "被问「测过真公告吗」再切恒兴新材——不要主动抢讲。",
        "（屏幕切到工作台，不要停在 PPT 上讲架构。）",
    ])
    _notes(s, "切工作台三栏。流程案固态电池；真公告被问再切恒兴。勿念闸族名单。")

    s = _new(prs)
    _title(s, "动作一：点开原文（约 70 秒）")
    _body(s, [
        "屏幕：证据与原文 → 点一条引文 → 「查看原文」",
        "",
        "口播要点：",
        "· 这不是模型转述，是从来源段落逐字摘出来的",
        "· 对不上原文的引文，进不了正式报告，会被打回",
        "· 这一步靠程序比对，不靠模型感觉像不像",
        "",
        "可选：指一下「少数意见」——吵输了还在，方便人再看一眼",
    ], size=16)
    _notes(s, "对照 docs/demo-script.md 动作一。综合意见每一点能点回结论。")

    s = _new(prs)
    _title(s, "动作二：改一个资料里没有的数字（约 80 秒）")
    _body(s, [
        "屏幕：在结论里加上「良率已达 99%」→ 保存改写",
        "",
        "同时发生两件事：",
        "1. 机械核对亮黄灯——原文里找不到这个数",
        "2. 顶部综合意见跟着改后的结论重写",
        "",
        "人可以改；编出来的数字不会假装没发生；操作记进日志。",
        "演示完可点「撤销改动」（也入链）。",
    ], size=16)
    _notes(s, "被施压改结论时：第一句说「可以，我们走修改流程」，不要偷改文件。")

    s = _new(prs)
    _title(s, "动作三：改一字即红（约 50 秒）")
    _body(s, [
        "屏幕：核验账本（绿）→ 演示篡改（红）",
        "",
        "口播要点：",
        "· 只在内存副本里改一处，磁盘文件不动",
        "· 诚实的是绿，改过的立刻红",
        "· 研判日志是账本，不是随便能改的文本",
        "",
        "收束三句：挂得住原文 · 人改得明 · 账本改一字就红",
    ], size=16)
    _notes(s, "收束后停。编队/MCP/置信度定义等评委问再答。")

    s = _new(prs)
    _title(s, "可选：被问真公告 → 切恒兴（约 40 秒）")
    _body(s, [
        "屏幕：案件「恒兴新材-评测」→ 已有 Q8 结果（或点开历史）",
        "议题：2024 半年度营收、归母净利、同比变动各是多少？",
        "",
        "指三下：",
        "· 综合意见标题三数齐：356,769,612.82 / 31,831,900.82 / -28.05%",
        "· 证据与原文 → 半年度报告摘要表行",
        "· 核验账本仍绿",
        "",
        "口播一句：固态电池是流程；恒兴是 8 份巨潮公告、10 问约 9.5/10。",
    ], size=15)
    _notes(s, "对照 eval/hengxing-defense.md。勿展开 Docling / 槽位调试史。")

    s = _new(prs)
    _title(s, "评委常问（备答，不要主动展开）")
    _body(s, [
        "和群聊 Agent 有何不同？→ 独立写、再裁决、输的留着、假引文打回",
        "0.73 是什么？→ 规则分不是概率；能拆开看、改完能重算",
        "测过真公告吗？→ 恒兴 8 份公告 / 10 问约 9.5/10；现场可切看 Q8",
        "断网 / 无 Key？→ 离线演示桩仍可完整跑通流程",
        "图片没读？→ 披露「没读」，假装读过比漏掉更贵",
        "怎么接到现有系统？→ MCP：读报告、验链、结论按契约入卷",
        "换题目还灵吗？→ 换资料夹 + 换议题；编队可用 roster 配置",
    ], size=14)
    _notes(s, "本页仅自己看。现场以工作台 live 为主，PPT 翻到动作页提醒节奏即可。")

    s = _new(prs)
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.5), Inches(2.2))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "现场主画面 = 工作台"
    _set_run(r, size=32, bold=True, color=ACCENT)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "PPT 只负责节奏；证据在屏幕上，不在幻灯片里"
    _set_run(r, size=16, color=DIM)

    path = _OUT / "veritas-现场讲解.pptx"
    prs.save(path)
    return path


def main():
    a = build_submission()
    b = build_onsite()
    _DESKTOP.mkdir(parents=True, exist_ok=True)
    for p in (a, b):
        dest = _DESKTOP / p.name
        shutil.copy2(p, dest)
        print(f"OK {p} -> {dest}")


if __name__ == "__main__":
    main()
